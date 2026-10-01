# -*- coding: utf-8 -*-
"""Checklist(Basshu) 를 자동 기입에 맞춘 구조로 재구성한 제안본을 만든다."""
import sys, io, os, re, json, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl, warnings; warnings.filterwarnings("ignore")
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

SRC, MAP, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
RELEASE = "2026.09.23 GME App 7.21.0 (New App)"

ABBR = [
    ("Registration", "REG"), ("Initial Screen", "INI"), ("Login", "LGN"), ("HOME", "HOM"),
    ("MENU", "MNU"), ("MY QR", "MQR"), ("Overseas Transfer", "OVS"), ("Domestic Transfer", "DOM"),
    ("Today's Rate", "RTE"), ("Domestic Topup", "DTU"), ("International Topup", "ITU"),
    ("Bill payment", "BIL"), ("Receive Overseas Funds", "RCV"), ("Deposit", "DEP"),
    ("Loan / Pesonal Finance", "LON"), ("ATM Withdraw", "ATM"), ("GME Wallet Statement", "WST"),
    ("Telecom", "TEL"), ("Issue Certificate", "CRT"), ("Booking", "BKG"), ("Card", "CRD"),
    ("Event", "EVT"), ("Profile", "PRF"),
]
ABBR_MAP = dict(ABBR)

YAML_FIX = {
    "04_Menu_History": "04_Menu_History_old.yaml",
    "06_Registration_All": "06_Registration_old.yaml",
    "02_02_Overseas_Schedule": "02_02_Overseas_Schedule_old.yaml",
    "03_Domestic": "03_Domestic_old.yaml",
    "08_MyQR": "08_MyQR_old.yaml",
    "05_InitialScreen": "05_InitialScreen_old.yaml",
    "01_01_Login_First": "01_01_Login_Success_old.yaml",
    "01_03_Login_Wrong simple password": "01_02_Login_Wrong_SimplePassword_old.yaml",
    "07_TodaysRate": "07_TodaysRate_old.yaml",
}
AREA_YAML = {
    "HOME": "09_Home_old.yaml",
    "Overseas Transfer": "02_01_Overseas_Sendnow_old.yaml",
    "Domestic Topup": "10_DomesticTopup_old.yaml",
    "International Topup": "11_IntlTopup_old.yaml",
    "Bill payment": "12_BillPayment_old.yaml",
    "Receive Overseas Funds": "13_ReceiveOverseas_old.yaml",
    "Deposit": "14_Deposit_old.yaml",
    "Loan / Pesonal Finance": "15_Loan_old.yaml",
    "ATM Withdraw": "16_ATMWithdraw_old.yaml",
    "GME Wallet Statement": "17_WalletStatement_old.yaml",
    "Telecom": "18_Telecom_old.yaml",
    "Issue Certificate": "19_IssueCertificate_old.yaml",
    "Booking": "20_Booking_old.yaml",
    "Card": "21_Card_old.yaml",
    "Event": "22_Event_old.yaml",
    "Profile": "25_Profile_old.yaml",
    "MENU": "26_Menu_old.yaml",
    "Registration": "06_Registration_old.yaml",
    "Initial Screen": "05_InitialScreen_old.yaml",
    "Login": "01_01_Login_Success_old.yaml",
    "Domestic Transfer": "03_Domestic_old.yaml",
    "Today's Rate": "07_TodaysRate_old.yaml",
    "MY QR": "08_MyQR_old.yaml",
}
TYPOS = {
    "Pesonal": "Personal", "Biometirc": "Biometric", "Authenication": "Authentication",
    "Balanece": "Balance", "Andorid": "Android", "Bamk": "Bank", "Forreigner": "Foreigner",
    "Carrirer": "Carrier", "Condtions": "Conditions", "KTFC": "KFTC",
    "Registeration": "Registration", "Requsest": "Request", "Procced": "Proceed",
    "Countinue": "Continue", "Depposit": "Deposit", "Shrare": "Share", "Eidt": "Edit",
    "HIstory": "History", "Histrory": "History", "occured": "occurred",
}

ws = openpyxl.load_workbook(SRC)["Checklist(Basshu)"]
archive = {x["row"]: x for x in json.load(open(MAP, encoding="utf-8"))}

rows = []
cur = {3: None, 4: None, 5: None}
for r in range(11, 424):
    f = {}
    for c in (3, 4, 5):
        v = ws.cell(r, c).value
        if v not in (None, ""):
            cur[c] = re.sub(r"\s+", " ", str(v)).strip()
            f[c] = False
        else:
            f[c] = True
    rows.append(dict(src=r, pri=ws.cell(r, 2).value,
                     d1=cur[3], d2=cur[4], d3=cur[5],
                     chk=ws.cell(r, 6).value or "", cmt=ws.cell(r, 10).value or "",
                     fill=dict(f)))

seq = collections.Counter()
for d in rows:
    code = ABBR_MAP[d["d1"]]
    seq[code] += 1
    d["id"] = "%s-%03d" % (code, seq[code])

dup = collections.defaultdict(list)
for d in rows:
    dup[re.sub(r"\s+", "", str(d["chk"]))].append(d["id"])
dupset = {k: v for k, v in dup.items() if len(v) > 1}

grp = collections.defaultdict(list)
for d in rows:
    grp[(d["d1"], d["d2"], d["d3"])].append(d)
pri_odd = set()
for g, members in grp.items():
    if len(members) < 3:
        continue
    cnt = collections.Counter(m["pri"] for m in members)
    if len(cnt) < 2:
        continue
    top, n = cnt.most_common(1)[0]
    if n >= len(members) - 1:
        for m in members:
            if m["pri"] != top:
                m["_pri_maj"] = top
                pri_odd.add(m["id"])

stats = collections.Counter()
for d in rows:
    memo, plat = [], ""
    t = str(d["chk"])
    if t.count("확인") >= 2 or ("\n" in t and t.count("확인") >= 1):
        memo.append("분리검토: 한 행에 확인 항목 2개 이상")
        stats["분리검토"] += 1
    key = re.sub(r"\s+", "", t)
    if key in dupset:
        memo.append("중복문구: " + ", ".join(x for x in dupset[key] if x != d["id"]) + " 와 동일")
        stats["중복문구"] += 1
    c = str(d["cmt"])
    if c:
        has_ios = re.search(r"ios", c, re.I) is not None
        has_aos = re.search(r"aos|\(And\)|Andorid|Android", c) is not None
        if has_ios and not has_aos:
            plat = "iOS"
            memo.append("플랫폼: Comment 가 iOS 만 언급")
            stats["플랫폼"] += 1
        elif has_aos and not has_ios:
            plat = "Android"
            memo.append("플랫폼: Comment 가 Android 만 언급")
            stats["플랫폼"] += 1
    if d["id"] in pri_odd:
        memo.append("Priority 확인: 같은 그룹은 %s 인데 이 행만 %s" % (d["_pri_maj"], d["pri"]))
        stats["Priority"] += 1
    own = [d["chk"], d["cmt"]] + [d[k] for k, col in (("d1", 3), ("d2", 4), ("d3", 5))
                                  if not d["fill"][col]]
    joined = " ".join(str(x) for x in own if x)
    found = [w for w in TYPOS if re.search(r"\b%s\b" % re.escape(w), joined)]
    if found:
        memo.append("오타후보: " + ", ".join("%s>%s" % (w, TYPOS[w]) for w in found))
        stats["오타"] += 1
    d["memo"] = " / ".join(memo)
    d["plat"] = plat

for d in rows:
    a = archive.get(d["src"])
    if a:
        d["yaml"] = "(제안) " + YAML_FIX.get(a["yaml"], a["yaml"] + ".yaml")
        d["auto"] = "예정"
        d["yaml_src"] = "검증된 과거 매핑"
    elif d["d1"] in AREA_YAML:
        d["yaml"] = "(제안) " + AREA_YAML[d["d1"]]
        d["auto"] = "예정"
        d["yaml_src"] = "영역 기준 추정"
    else:
        d["yaml"] = ""
        d["auto"] = ""
        d["yaml_src"] = ""

wb = openpyxl.Workbook()
sh = wb.active
sh.title = "Checklist(Basshu)"

HDR = ["ID", "Priority", "1Depth", "2Depth", "3Depth", "Checklist",
       "대상 플랫폼", "사전조건·테스트데이터", "기대 결과", "자동화 여부",
       "yaml 파일", "검토 메모", "Comment",
       RELEASE + " | iOS", RELEASE + " | Android"]
WID = [11, 10, 20, 24, 26, 58, 12, 26, 30, 11, 34, 52, 46, 22, 22]

hfill = PatternFill("solid", fgColor="1F4E78")
hfont = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
thin = Side(style="thin", color="BFBFBF")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
gray = Font(name="맑은 고딕", size=10, color="A6A6A6")
norm = Font(name="맑은 고딕", size=10)
memofont = Font(name="맑은 고딕", size=9, color="C55A11")

for i, h in enumerate(HDR, 1):
    c = sh.cell(1, i, h)
    c.fill, c.font, c.border = hfill, hfont, bd
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sh.column_dimensions[L(i)].width = WID[i - 1]
sh.row_dimensions[1].height = 34

for i, d in enumerate(rows, 2):
    vals = [d["id"], d["pri"], d["d1"], d["d2"], d["d3"], d["chk"], d["plat"],
            "", "", d["auto"], d["yaml"], d["memo"], d["cmt"], "", ""]
    for j, v in enumerate(vals, 1):
        c = sh.cell(i, j, v)
        c.border = bd
        c.alignment = Alignment(vertical="center", wrap_text=(j in (6, 12, 13)))
        if j in (3, 4, 5) and d["fill"][j]:
            c.font = gray
        elif j == 12:
            c.font = memofont
        else:
            c.font = norm
    sh.cell(i, 1).font = Font(name="Consolas", size=10, bold=True)

sh.freeze_panes = "G2"
sh.auto_filter.ref = "A1:%s%d" % (L(len(HDR)), len(rows) + 1)

dv_plat = DataValidation(type="list", formula1='"iOS,Android,공통"', allow_blank=True)
dv_auto = DataValidation(type="list", formula1='"Y,N,예정,불가"', allow_blank=True)
dv_res = DataValidation(type="list", formula1='"Pass,Fail,N/A,Not Test,Blocked"', allow_blank=True)
dv_res_b = DataValidation(type="list", formula1='"Pass,Fail,N/A,Not Test,Blocked"', allow_blank=True)
for dv, col in ((dv_plat, "G"), (dv_auto, "J"), (dv_res, "N"), (dv_res_b, "O")):
    sh.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col, col, len(rows) + 1))

rng = "N2:O%d" % (len(rows) + 1)
for val, bgc, fgc in (("Pass", "C6EFCE", "006100"), ("Fail", "FFC7CE", "9C0006"),
                      ("Blocked", "FFEB9C", "9C6500"), ("Not Test", "D9D9D9", "808080"),
                      ("N/A", "EDEDED", "808080")):
    sh.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"%s"' % val],
        fill=PatternFill("solid", fgColor=bgc), font=Font(color=fgc)))

cv = wb.create_sheet("코드값")
for j, w in enumerate([18, 30, 46, 16], 1):
    cv.column_dimensions[L(j)].width = w


def block(r0, title, header, data):
    c = cv.cell(r0, 1, title)
    c.font = Font(name="맑은 고딕", bold=True, size=12, color="1F4E78")
    for j, h in enumerate(header, 1):
        x = cv.cell(r0 + 1, j, h)
        x.fill, x.font, x.border = hfill, hfont, bd
    for i, row in enumerate(data, r0 + 2):
        for j, v in enumerate(row, 1):
            x = cv.cell(i, j, v)
            x.border = bd
            x.font = norm
            x.alignment = Alignment(vertical="center", wrap_text=True)
    return r0 + 2 + len(data) + 1


n = block(1, "결과값", ["값", "의미", "쓰는 때"], [
    ["Pass", "통과", "기대 결과대로 동작"],
    ["Fail", "실패", "기대 결과와 다르게 동작 (결함)"],
    ["N/A", "해당 없음", "이 플랫폼·계정에는 존재하지 않는 기능"],
    ["Not Test", "미수행", "이번 회차에 돌리지 않음"],
    ["Blocked", "수행 불가", "선행 조건·데이터·외부 인증 때문에 막힘"],
])
n = block(n, "자동화 여부", ["값", "의미"], [
    ["Y", "자동화되어 있고 스위트에서 돌고 있음"],
    ["N", "자동화하지 않음 (수동 유지)"],
    ["예정", "자동화 대상이지만 아직 작성·검증 전"],
    ["불가", "외부 인증·실결제 등으로 자동화가 불가능"],
])
n = block(n, "대상 플랫폼", ["값", "의미"], [
    ["iOS", "iOS 에만 해당"],
    ["Android", "Android 에만 해당"],
    ["공통", "두 플랫폼 모두 해당 (비워두면 공통으로 본다)"],
])
n = block(n, "환경", ["값", "의미"], [
    ["Stage", "스테이지 서버"],
    ["Livetest Old app", "라이브테스트 · 구 UI 빌드"],
    ["Livetest New app", "라이브테스트 · 개편 UI 빌드"],
    ["Live Beta", "운영 베타"],
])
n = block(n, "ID 영역 약어", ["약어", "1Depth", "현재 개수"],
          [[code, name, sum(1 for d in rows if d["d1"] == name)] for name, code in ABBR])

c = cv.cell(n, 1, "ID 규칙")
c.font = Font(name="맑은 고딕", bold=True, size=12, color="1F4E78")
for i, t in enumerate([
    "1. 형식은 <영역약어>-<3자리 일련번호> 이다. 예: REG-001",
    "2. 한 번 부여한 ID 는 바꾸지 않는다. 문구가 수정돼도 같은 항목이면 ID 를 유지한다.",
    "3. 삭제된 항목의 번호는 재사용하지 않는다. 그 번호는 영구히 비워 둔다.",
    "4. 항목이 새로 생기면 그 영역의 마지막 번호 +1 을 쓴다. 중간에 끼워 넣지 않는다.",
    "5. 한 항목을 둘로 쪼개면, 원래 ID 는 그대로 두고 나머지 하나에 새 번호를 준다.",
    "6. 영역(1Depth)이 새로 생기면 위 표에 약어를 먼저 등록하고 쓴다. 약어는 영문 3글자.",
    "7. 자동화 스크립트는 행 번호가 아니라 ID 로 행을 찾는다. 행을 넣고 빼도 안전하다.",
], 1):
    x = cv.cell(n + i, 1, t)
    x.font = norm
    cv.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=4)

ar = wb.create_sheet("자동화결과")
AH = ["실행일시", "릴리스", "환경", "플랫폼", "빌드", "기기", "ID", "결과", "yaml 파일", "비고"]
AW = [20, 34, 20, 11, 16, 20, 12, 12, 34, 40]
for j, h in enumerate(AH, 1):
    c = ar.cell(1, j, h)
    c.fill, c.font, c.border = hfill, hfont, bd
    c.alignment = Alignment(horizontal="center", vertical="center")
    ar.column_dimensions[L(j)].width = AW[j - 1]
ar.freeze_panes = "A2"
ar.auto_filter.ref = "A1:J1"
dv_env = DataValidation(type="list", formula1='"Stage,Livetest Old app,Livetest New app,Live Beta"', allow_blank=True)
dv_pf = DataValidation(type="list", formula1='"iOS,Android"', allow_blank=True)
dv_res2 = DataValidation(type="list", formula1='"Pass,Fail,N/A,Not Test,Blocked"', allow_blank=True)
for dv, col in ((dv_env, "C"), (dv_pf, "D"), (dv_res2, "H")):
    ar.add_data_validation(dv)
    dv.add("%s2:%s5000" % (col, col))

wb.save(OUT)

print("생성:", OUT)
print("행 %d (원본 11~423 = %d행)" % (len(rows), 423 - 11 + 1))
print("ID 영역 %d개, 중복 ID %d" % (len(seq), len(rows) - len(set(d["id"] for d in rows))))
print("\n검토 메모 집계:")
for k, v in stats.most_common():
    print("   %-10s %d" % (k, v))
print("   메모가 붙은 행 %d / %d" % (sum(1 for d in rows if d["memo"]), len(rows)))
print("\nyaml 제안:")
for k, v in collections.Counter(d["yaml_src"] for d in rows).most_common():
    print("   %-14s %d" % (k or "(없음)", v))

json.dump([{k: d[k] for k in ("id", "src", "pri", "d1", "d2", "d3", "chk", "cmt",
                              "plat", "auto", "yaml", "yaml_src", "memo")} for d in rows],
          open(os.path.join(os.path.dirname(OUT), "rows_new.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
